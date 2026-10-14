# -*- coding: utf-8 -*-
"""在已汇报的《Openclaw-RL论文复现及迁移内容分享》第 5 页之后插入 MetaClaw 进展 4 页。

- 输入（只读，不改动）：Openclaw-RL论文复现及迁移内容分享.pptx（2026-09-18 已汇报版）
- 输出：Openclaw-RL论文复现及迁移内容分享_MetaClaw进展.pptx
- 表格数字全部由 metaclaw_progress_data/ 的逐日 CSV 现算（compute_windows.agg），不手抄
- 页面内容与口径见 metaclaw进展_新增四页方案.md

用法：python build_metaclaw_progress_slides.py
"""
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "metaclaw_progress_data"))
import compute_windows as cw  # noqa: E402

SRC = os.path.join(HERE, "Openclaw-RL论文复现及迁移内容分享.pptx")
OUT = os.path.join(HERE, "Openclaw-RL论文复现及迁移内容分享_MetaClaw进展.pptx")
INSERT_AFTER = 5            # 插在第 5 页之后
TEMPLATE_SLIDE = 5          # 版式取自第 5 页

LATIN_TXT, EA = "Segoe UI", "微软雅黑"
LATIN_NUM = "Times New Roman"    # 第 4、5 页表格图里数字是衬线体
INK = "262626"
GREY = "7F7F7F"
BASE_GREY = "595959"
RED = "B02A2A"                   # 第 5 页 ×1388 用的红
RULE = "404040"
THIN = "BFBFBF"
NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # No Style, No Grid

RUNS = cw.load()

ROW_LABEL = {
    "B0": "零训练基线 B0", "B0'": "零训练基线 B0'",
    "H0": "混合（对照）", "OPD": "纯 OPD（去掉 RL）", "RL": "纯 RL（去掉 OPD）",
    "ENT": "+ 熵 0.1", "KL": "+ KL 0.1",
    "MASK": "纯 OPD + 遮罩", "ONP": "混合 + 同策略",
}
FOOTNOTE = ("Acc / Compl（%）。Compl 与第 5 页同口径：file_check 通过数 / 当天全部题数；窗口按题数加总。"
            "“—” = 该趟未跑到这一段（没跑到的天不计 0）。红字“复读 xx%” = 该段 copies≥8 的 turn 占比 ≥ 25%；"
            "红色数值 = Acc < 10%。B0' 为第二趟零训练基线，作噪声参照。")


# ---------------------------------------------------------------- 文本工具
def _rpr(run, sz, bold=False, color=INK, latin=LATIN_TXT):
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "zh-CN")
    rPr.set("altLang", "en-US")
    rPr.set("sz", str(int(sz * 100)))
    rPr.set("b", "1" if bold else "0")
    rPr.set("noProof", "1")
    for tag in ("a:solidFill", "a:latin", "a:ea"):
        for el in rPr.findall(qn(tag)):
            rPr.remove(el)
    fill = etree.SubElement(rPr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr")).set("val", color)
    etree.SubElement(rPr, qn("a:latin")).set("typeface", latin)
    etree.SubElement(rPr, qn("a:ea")).set("typeface", EA)


def _ppr(p, spc_aft_pt=5, ln_pct=112, algn=None):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    if algn:
        pPr.set("algn", algn)
    for child in list(pPr):
        pPr.remove(child)
    ln = etree.SubElement(pPr, qn("a:lnSpc"))
    etree.SubElement(ln, qn("a:spcPct")).set("val", str(ln_pct * 1000))
    sb = etree.SubElement(pPr, qn("a:spcBef"))
    etree.SubElement(sb, qn("a:spcPts")).set("val", "0")
    sa = etree.SubElement(pPr, qn("a:spcAft"))
    etree.SubElement(sa, qn("a:spcPts")).set("val", str(int(spc_aft_pt * 100)))
    etree.SubElement(pPr, qn("a:buNone"))


def write_paras(tf, paras, ln_pct=112):
    """paras: [(spc_aft_pt, [(text, sz, bold, color), ...]), ...]"""
    tf.clear()
    for i, (aft, runs) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _ppr(p, aft, ln_pct)
        for text, sz, bold, color in runs:
            r = p.add_run()
            r.text = text
            _rpr(r, sz, bold, color)


def body(slide, x, y, w, h, heading, items):
    """用第 5 页同款内容占位符写「标题 + ①②③」要点。"""
    ph = slide.placeholders[1]
    ph.left, ph.top, ph.width, ph.height = Inches(x), Inches(y), Inches(w), Inches(h)
    paras = [(6, [(heading, 18, True, INK)])] if heading else []
    for i, it in enumerate(items):
        runs = it if isinstance(it, list) else [(it, 14, False, INK)]
        paras.append((0 if i == len(items) - 1 else 5, runs))
    write_paras(ph.text_frame, paras)
    ph.text_frame.word_wrap = True


def textbox(slide, x, y, w, h, runs, sz=11, color=GREY, bold=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    if isinstance(runs, str):
        runs = [(runs, sz, bold, color)]
    write_paras(tf, [(0, runs)], ln_pct=105)
    return tb


def title(slide, text):
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = Emu(484632), Emu(256032), Emu(11155680), Emu(822960)
    tf = t.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p._p.get_or_add_pPr().set("algn", "l")
    r = p.add_run()
    r.text = text
    _rpr(r, 26, True, INK)


def panel(slide, x, y, w, h):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shp.line.fill.background()
    shp.shadow.inherit = False
    shp.name = "Panel"
    return shp


# ---------------------------------------------------------------- 表格工具
def _ln(tag, w_pt=None, color=None):
    ln = etree.Element(qn(tag))
    if w_pt is None:
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))
    else:
        ln.set("w", str(int(w_pt * 12700)))
        fill = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", color)
    return ln


def style_cell(cell, top=None, bottom=None, fill="FFFFFF"):
    """top/bottom: (宽 pt, 颜色) 或 None。按 schema 顺序 lnL lnR lnT lnB fill 重写 tcPr。"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        tcPr.remove(child)
    tcPr.set("anchor", "ctr")
    for side in ("L", "R", "T", "B"):
        spec = {"T": top, "B": bottom}.get(side)
        tcPr.append(_ln(f"a:ln{side}", *(spec if spec else (None, None))))
    sf = etree.SubElement(tcPr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", fill)
    for side in ("left", "right"):
        setattr(cell, f"margin_{side}", Inches(0.06))
    cell.margin_top = cell.margin_bottom = Inches(0.03)


def fill_cell(cell, lines, align="ctr"):
    """lines: [[(text, sz, bold, color, latin), ...], ...]，每个子列表一段。"""
    tf = cell.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, runs in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _ppr(p, 0, 100, algn=align)
        for text, sz, bold, color, latin in runs:
            r = p.add_run()
            r.text = text
            _rpr(r, sz, bold, color, latin)


def make_table(slide, x, y, col_w, row_h, header, rows):
    """header: [str]；rows: [[cell_lines, ...], ...]，cell_lines 见 fill_cell。"""
    n_r, n_c = len(rows) + 1, len(header)
    gf = slide.shapes.add_table(n_r, n_c, Inches(x), Inches(y),
                                Inches(sum(col_w)), Inches(row_h[0] + sum(row_h[1:])))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = NO_STYLE
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)
    for i, h in enumerate(row_h):
        tbl.rows[i].height = Inches(h)
    for j, h in enumerate(header):
        c = tbl.cell(0, j)
        style_cell(c, top=(1.25, RULE), bottom=(0.75, RULE))
        fill_cell(c, [[(h, 12, True, INK, LATIN_TXT)]], align="l" if j == 0 else "ctr")
    for i, row in enumerate(rows, start=1):
        last = i == n_r - 1
        for j, lines in enumerate(row):
            c = tbl.cell(i, j)
            style_cell(c, bottom=(1.25, RULE) if last else (0.5, THIN))
            fill_cell(c, lines, align="l" if j == 0 else "ctr")
    return gf


def label_cell(tag):
    color = BASE_GREY if tag.startswith("B0") else INK
    return [[(ROW_LABEL[tag], 12, not tag.startswith("B0"), color, LATIN_TXT)]]


def window_cell(tag, a, b):
    x = cw.agg(RUNS, tag, a, b)
    if x is None:
        return [[("—", 12, False, GREY, LATIN_NUM)]]
    base = tag.startswith("B0")
    collapsed = x["acc"] < 10
    num_color = RED if collapsed else (BASE_GREY if base else INK)
    lines = [[(f"{x['acc']:.1f} / {x['pa']:.1f}", 12.5, collapsed, num_color, LATIN_NUM)]]
    if x["rep"] == x["rep"] and x["rep"] >= 100 * cw.ONSET_THR:
        lines.append([(f"复读 {x['rep']:.0f}%", 9.5, False, RED, LATIN_TXT)])
    return lines


def window_header(a, b):
    return f"day{a:02d}" if a == b else f"day{a:02d}–{b:02d}"


def onset_text(tag):
    if tag.startswith("B0") or tag == "RL":
        return "无"
    d = cw.onset(RUNS, tag)
    if d is None:
        return f"到 day{cw.last_full(RUNS, tag):02d} 无"
    return f"day{d:02d}"


def dropped_total(tag):
    return sum(int(r["dropped"]) for r in RUNS[tag].values())


def onset_cell(tag):
    base = tag.startswith("B0")
    color = BASE_GREY if base else INK
    lines = [[(onset_text(tag), 11, False, color, LATIN_TXT)]]
    if tag == "RL":   # 纯 RL 不复读 thinking，崩法单独说明
        lines.append([("day12 起 Acc 归零", 9.5, False, RED, LATIN_TXT)])
    lines.append([(f"DROPPED {dropped_total(tag)}", 9.5, False, GREY, LATIN_TXT)])
    return lines


# ---------------------------------------------------------------- 页面
def comparison_slide(slide, title_text, caption, tags, wins, bullets_heading, bullets):
    title(slide, title_text)
    px, py, pw = 0.53, 1.18, 12.27
    label_w, onset_w = 2.15, 1.85
    seg_w = (pw - 0.3 - label_w - onset_w) / len(wins)
    col_w = [label_w] + [seg_w] * len(wins) + [onset_w]
    row_h = [0.38] + [0.52] * len(tags)
    tbl_h = sum(row_h)
    ph = 0.46 + tbl_h + 0.58
    panel(slide, px, py, pw, ph)
    textbox(slide, px + 0.15, py + 0.10, pw - 0.3, 0.3, caption, sz=13, color=INK, bold=True)
    header = ["实验"] + [window_header(a, b) for a, b in wins] + ["成批复读首日"]
    rows = [[label_cell(t)] + [window_cell(t, a, b) for a, b in wins] + [onset_cell(t)] for t in tags]
    make_table(slide, px + 0.15, py + 0.46, col_w, row_h, header, rows)
    textbox(slide, px + 0.15, py + 0.46 + tbl_h + 0.07, pw - 0.3, 0.5, FOOTNOTE, sz=9.5)
    by = py + ph + 0.14
    body(slide, 0.53, by, 12.27, 7.35 - by, bullets_heading, bullets)


def mechanism_half(slide, x, w, caption, sub, tags, wins, heading, bullets):
    py = 1.18
    label_w = 1.75
    seg_w = (w - 0.3 - label_w) / len(wins)
    col_w = [label_w] + [seg_w] * len(wins)
    row_h = [0.40] + [0.56] * len(tags)
    tbl_h = sum(row_h)
    ph = 0.78 + tbl_h + 0.14
    panel(slide, x, py, w, ph)
    textbox(slide, x + 0.15, py + 0.08, w - 0.3, 0.3, caption, sz=13, color=INK, bold=True)
    textbox(slide, x + 0.15, py + 0.42, w - 0.3, 0.3, sub, sz=10.5, color=GREY)
    header = ["实验"] + [window_header(a, b) for a, b in wins]
    rows = [[label_cell(t)] + [window_cell(t, a, b) for a, b in wins] for t in tags]
    make_table(slide, x + 0.15, py + 0.78, col_w, row_h, header, rows)
    by = py + ph + 0.14
    body_box = slide.shapes.add_textbox(Inches(x), Inches(by), Inches(w), Inches(6.62 - by))
    tf = body_box.text_frame
    tf.word_wrap = True
    paras = [(6, [(heading, 16, True, INK)])]
    for i, it in enumerate(bullets):
        runs = it if isinstance(it, list) else [(it, 13, False, INK)]
        paras.append((0 if i == len(bullets) - 1 else 5, runs))
    write_paras(tf, paras)
    return by


def build_slide6(s):
    tags, wins = cw.GROUPS["消融"]
    comparison_slide(
        s, "消融实验：RL 与 OPD 各自的作用",
        "对照：混合训练（第 5 页那趟）　｜　变量：去掉 RL / 去掉 OPD",
        tags, wins, "结论",
        ["① 零训练两趟都不崩（DROPPED = 0）：崩溃由训练引起，不是题集本身的性质。",
         "② 复读 thinking 只出现在开了 OPD 的训练里：纯 OPD day04 起病，比混合早 5 天"
         "（停在 day05 第 8 轮，当天已跑的 8 轮全部超时）。",
         "③ 纯 RL 不复读 thinking，但 day12 起 Acc 归零直到 day27：day13 起当天 main turn 从约 20 涨到 141，"
         "反复写同一文件（day18 一轮 write ×499），最后上下文撑爆、压缩超时。"])


def build_slide7(s):
    tags, wins = cw.GROUPS["正则"]
    comparison_slide(
        s, "正则实验：能否把崩溃挡住",
        "对照：混合训练（第 5 页那趟）　｜　变量：+ 熵正则 0.1 / + KL 正则 0.1（锚定初始模型）",
        tags, wins, "结论",
        ["① 熵正则：熵确实被抬起来了（entropy 0.157 → 0.638），起病反而从 day09 提前到 day06。",
         "② KL：起病推迟到 day18。day09–17 对照和熵趟都已崩，KL 与基线基本持平（39.3 vs 41.8）；"
         "但起病后 3 天内复读占比升到 71%，最终没有挡住。",
         "③ day01–05 各趟在 67.6–75.6 之间，与两趟零训练基线自身的差距（67.6 vs 74.4）同量级，不能看作收益。"])


def build_slide8(s):
    title(s, "机制验证实验")
    w = 6.05
    tags, wins = cw.GROUPS["机制·遮罩"]
    mechanism_half(
        s, 0.53, w, "假设 H-e：OPD 压低了“停止思考”",
        "对照：纯 OPD　｜　实验：遮住 OPD 对停止决策及随后动作的作用",
        tags, wins, "结果：否定",
        ["① 遮罩生效：每步 OPD 仍作用的位置占比 0.83–0.98，7 步均在预设区间 [0.8, 1.0) 内。",
         "② 两趟都是 day04 起病，遮罩趟 day05 起复读占比 97%：一天没推迟，问题出在 OPD 作用于 thinking 内容的部分。"])
    tags, wins = cw.GROUPS["机制·同策略"]
    mechanism_half(
        s, 0.53 + w + 0.17, w, "假设 H-f：训练用了旧权重生成的样本",
        "对照：混合训练　｜　实验：只用当前权重生成的样本训练",
        tags, wins, "结果：未定论",
        ["① 到 day09 没有成批复读（对照 day09 为 29%）；但 day09 结束时只训练到 step 7，对照起病时已到 step 11。",
         "② 出现了别的失效：day08 一个 95 turn 的超长轮次（不是复读 thinking），当天 Acc 0；day09 Compl 0。"])
    textbox(s, 0.53, 6.72, 12.27, 0.5, FOOTNOTE, sz=9.5)
    # 这一页两半自带标题，不用内容占位符
    s.shapes._spTree.remove(s.placeholders[1]._element)


def build_slide9(s):
    title(s, "总结")
    px, py, pw = 0.53, 1.18, 12.27
    header = ["实验", "对照", "成批复读首日（对照 → 实验）", "结论"]
    col_w = [2.6, 1.9, 3.5, 3.97]
    data = [
        ("去掉 RL（纯 OPD）", "混合", "day09 → day04", "复读 thinking 只在开了 OPD 时出现"),
        ("去掉 OPD（纯 RL）", "混合", "day09 → 无复读；day12 起 Acc 归零", "换了一种崩法：反复 write、上下文撑爆"),
        ("+ 熵 0.1", "混合", "day09 → day06", "加速崩溃"),
        ("+ KL 0.1", "混合", "day09 → day18", "推迟 9 天，挡不住"),
        ("遮住停止决策（H-e）", "纯 OPD", "day04 → day04", "一天没推迟，假设否定"),
        ("只用当前权重样本（H-f）", "混合", "day09 → 到 day09 无复读", "步数没对齐、出现别的失效，未定论"),
    ]
    rows = []
    for name, ctrl, ons, concl in data:
        rows.append([[[(name, 12, True, INK, LATIN_TXT)]],
                     [[(ctrl, 12, False, INK, LATIN_TXT)]],
                     [[(ons, 12, False, INK, LATIN_TXT)]],
                     [[(concl, 12, False, INK, LATIN_TXT)]]])
    row_h = [0.40] + [0.44] * len(rows)
    ph = 0.2 + sum(row_h) + 0.2
    panel(s, px, py, pw, ph)
    make_table(s, px + 0.15, py + 0.2, col_w, row_h, header, rows)
    by = py + ph + 0.18
    body(s, 0.53, by, 12.27, 7.35 - by, None,
         [[("① ", 14, False, INK), ("零训练基线跑满 30 天不崩；目前所有训练配置都在 30 天内崩溃，没有一趟超过零训练基线。", 14, False, INK)],
          [("② ", 14, False, INK), ("崩法有两种：开了 OPD → 复读 thinking → idle timeout；纯 RL → 反复 write → 上下文撑爆、压缩超时。", 14, False, INK)],
          [("③ ", 14, False, INK), ("正则只能改变起病时间：熵提前 3 天，KL 推迟 9 天；两个机制假设中 H-e 已否定，H-f 尚无结论。", 14, False, INK)],
          [("④ ", 14, False, INK), ("口径：健康期 ±10pt 以内的差异在两趟零训练基线的差距范围内；训练中途停止的趟，没跑到的天不计入。", 14, False, INK)]])


def main():
    prs = Presentation(SRC)
    n0 = len(prs.slides)
    layout = prs.slides[TEMPLATE_SLIDE - 1].slide_layout
    builders = [build_slide6, build_slide7, build_slide8, build_slide9]
    for b in builders:
        b(prs.slides.add_slide(layout))
    # 新页移到第 INSERT_AFTER 页之后
    lst = prs.slides._sldIdLst
    ids = list(lst)
    new = ids[n0:]
    for el in new:
        lst.remove(el)
    for k, el in enumerate(new):
        lst.insert(INSERT_AFTER + k, el)
    prs.save(OUT)
    print("wrote", OUT, "slides:", len(prs.slides))


if __name__ == "__main__":
    main()
