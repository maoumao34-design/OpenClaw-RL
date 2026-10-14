# -*- coding: utf-8 -*-
"""从 5 页原版出发，按用户新口径改第 4 页并插入第 5 页。前三页与 THANKS 不动。"""
from __future__ import annotations

import copy
import re
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(r"D:\MAO\Claude\openclaw-rl\report")
PPT = HERE / "Openclaw-RL复现及迁移内容分享.pptx"
PAPER_TBL = HERE / "table1_full.png"
OURS_TBL = HERE / "perq_k0_k6.png"

EA_FONT = "微软雅黑"
LATIN_FONT = "Segoe UI"
F = "C:/Windows/Fonts/msyh.ttc"
FB = "C:/Windows/Fonts/msyhbd.ttc"
SERIF = "C:/Windows/Fonts/times.ttf"
SERIF_BD = "C:/Windows/Fonts/timesbd.ttf"

INK = (34, 38, 46)
MUTED = (110, 118, 130)
HILITE = (176, 42, 42)
RULE = (0, 0, 0)
OURS_BG = (255, 244, 214)
BG = (255, 255, 255)


def style_run(run, size, bold=None, color=None):
    font = run.font
    font.size = size
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = RGBColor(*color)
    font.name = LATIN_FONT
    rPr = font._rPr
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        latin = rPr.find(qn("a:latin"))
        latin.addnext(ea) if latin is not None else rPr.append(ea)
    ea.set("typeface", EA_FONT)
    rPr.set("lang", "zh-CN")
    rPr.set("altLang", "en-US")
    rPr.set("noProof", "1")


def add_run(para, text, size, bold=None, color=None):
    run = para.add_run()
    run.text = text
    style_run(run, size, bold=bold, color=color)
    return run


def reset_bullet(para):
    pPr = para._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    for attr in ("marL", "indent"):
        pPr.attrib.pop(attr, None)


def set_no_bullet(para):
    pPr = para._p.get_or_add_pPr()
    pPr.attrib["marL"] = "0"
    pPr.attrib["indent"] = "0"
    if not pPr.findall(qn("a:buNone")):
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def set_rich_text(text_frame, items, head_size, body_size,
                  line_spacing=1.03, space_after=Pt(4)):
    text_frame.clear()
    text_frame.word_wrap = True
    for i, (level, raw) in enumerate(items):
        para = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        para.level = level
        reset_bullet(para)
        para.line_spacing = line_spacing
        para.space_before = Pt(0)
        para.space_after = space_after if level else Pt(4)
        size = head_size if level == 0 else body_size
        for seg in re.split(r"(\*\*.+?\*\*)", raw):
            if not seg:
                continue
            run = para.add_run()
            bold = seg.startswith("**") and seg.endswith("**")
            run.text = seg[2:-2] if bold else seg
            style_run(run, size, bold=True if bold else None)


def disable_proofing(slide):
    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes)
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    rPr = run.font._rPr
                    rPr.set("lang", "zh-CN")
                    rPr.set("altLang", "en-US")
                    rPr.set("noProof", "1")
    walk(slide.shapes)


def render_ours_table():
    """按题隔离 day 0 / day 6，带 Model 列。"""
    W, H = 1140, 280
    img = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(img)
    title = ImageFont.truetype(FB, 28)
    cn = ImageFont.truetype(F, 26)
    cnb = ImageFont.truetype(FB, 26)
    num = ImageFont.truetype(SERIF, 26)
    numb = ImageFont.truetype(SERIF_BD, 26)
    note = ImageFont.truetype(F, 22)

    d.text((12, 8), "本次按题隔离（同一 driver 路径）", font=title, fill=INK)
    d.line([(12, 48), (W - 12, 48)], fill=RULE, width=2)

    heads = [(12, "Model"), (300, "Condition"), (700, "Acc. (%)"), (920, "Compl. (%)")]
    for x, t in heads:
        d.text((x, 58), t, font=numb, fill=INK)
    d.line([(12, 92), (W - 12, 92)], fill=RULE, width=1)

    rows = [
        ("Qwen3-4B", "day 0  零训练基线", "34.4", "12.1", False),
        ("Qwen3-4B", "day 6  Hybrid RL", "37.3", "13.9", True),
    ]
    y = 104
    for model, cond, acc, compl, hi in rows:
        if hi:
            d.rectangle([8, y - 6, W - 8, y + 36], fill=OURS_BG)
        cond_font = cnb if hi else cn
        num_font = numb if hi else num
        color = HILITE if hi else INK
        d.text((12, y), model, font=cond_font, fill=color)
        d.text((300, y), cond, font=cond_font, fill=color)
        d.text((730, y), acc, font=num_font, fill=color)
        d.text((960, y), compl, font=num_font, fill=color)
        y += 42
    d.line([(12, y - 4), (W - 12, y - 4)], fill=RULE, width=2)
    d.text((12, y + 8), "相对提升 Acc +2.9pt / Compl +1.8pt",
           font=note, fill=MUTED)
    img.save(OURS_TBL)
    print("wrote", OURS_TBL.name, img.size)


def snapshot_slides(prs, n=3):
    return [etree.tostring(prs.slides[i]._element) for i in range(n)]


def copy_page_number(src_slide, dst_slide):
    existing = [
        sh for sh in dst_slide.placeholders
        if sh.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER
    ]
    if existing:
        return
    donor = next(
        (sh for sh in src_slide.placeholders
         if sh.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER),
        None,
    )
    if donor is None:
        return
    dst_slide.shapes._spTree.append(copy.deepcopy(donor._element))


def move_slide(prs, old_index, new_index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    el = slides[old_index]
    xml_slides.remove(el)
    xml_slides.insert(new_index, el)


def update_slide4(slide):
    title = next(sh for sh in slide.placeholders if sh.placeholder_format.idx == 0)
    assert "方法泛化方案" in "".join(r.text for r in title.text_frame.paragraphs[0].runs)

    # 只删左上旧 Table 1，右栏迁移图保留
    for shape in list(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.left.inches < 6:
            shape._element.getparent().remove(shape._element)

    left, top, width = 0.53, 1.08, 5.90
    paper = slide.shapes.add_picture(str(PAPER_TBL), Inches(left), Inches(top), width=Inches(width))
    y = top + paper.height.inches + 0.04

    cap1 = slide.shapes.add_textbox(Inches(left), Inches(y), Inches(width), Inches(0.28))
    tf = cap1.text_frame
    tf.word_wrap = True
    add_run(tf.paragraphs[0], "论文 Table 1　·　Part I / Part II 全表（Acc. / Compl.）", Pt(11), bold=True)
    y += 0.30

    ours = slide.shapes.add_picture(str(OURS_TBL), Inches(left), Inches(y), width=Inches(width))
    y += ours.height.inches + 0.04

    # 原说明框若还在就改内容并下移，否则新建
    caption = None
    for shape in slide.shapes:
        if shape.has_text_frame and not shape.is_placeholder:
            t = "".join(p.text for p in shape.text_frame.paragraphs)
            if "仅迁移 RL" in t or "数量级参照" in t:
                caption = shape
                break
    cap_text = "上表为本次按题隔离，不与论文 Table 1 数字直接对比"
    if caption is None:
        caption = slide.shapes.add_textbox(Inches(left), Inches(y), Inches(width), Inches(0.40))
    else:
        caption.left = Inches(left)
        caption.top = Inches(y)
        caption.width = Inches(width)
        caption.height = Inches(0.40)
    tf = caption.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.line_spacing = 1.08
    add_run(p, cap_text, Pt(10))
    y += 0.42

    body = next(sh for sh in slide.placeholders if sh.placeholder_format.idx == 1)
    body.left, body.width = Inches(0.53), Inches(5.90)
    body.top = Inches(min(y, 5.55))
    body.height = Inches(max(1.55, 7.32 - body.top.inches))
    set_rich_text(body.text_frame, [
        (0, "**① 任务背景**"),
        (1, "方法已在 Personal Agent 场景复现校准，换到 **MetaClaw-Bench**"
            "（30 天连续任务流、CLI agent）验证能否迁移"),
        (0, "**② 按题隔离：训练有效果**"),
        (1, "day 6 相对同路径 day 0：Acc +2.9pt / Compl +1.8pt"),
    ], Pt(13.5), Pt(11.5), line_spacing=1.03, space_after=Pt(4))
    disable_proofing(slide)


def fill_slide5(slide, donor):
    copy_page_number(donor, slide)
    title = next(sh for sh in slide.placeholders if sh.placeholder_format.idx == 0)
    title.left, title.top = Inches(0.53), Inches(0.28)
    title.width, title.height = Inches(12.20), Inches(0.90)
    tf = title.text_frame
    tf.clear()
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    add_run(tf.paragraphs[0], "按天隔离实验", Pt(26), bold=True)

    body = next(sh for sh in slide.placeholders if sh.placeholder_format.idx == 1)
    body.left, body.top = Inches(0.53), Inches(1.30)
    body.width, body.height = Inches(6.15), Inches(5.80)
    tf = body.text_frame
    tf.clear()
    tf.word_wrap = True

    def p(first=False):
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        para.level = 0
        set_no_bullet(para)
        return para

    p0 = p(True)
    p0.space_after = Pt(6)
    add_run(p0, "① 定版基线（day 0）", Pt(15), bold=True)

    p1 = p()
    p1.space_after = Pt(2)
    add_run(p1, "Acc.　　", Pt(12))
    add_run(p1, "41.1%", Pt(24), bold=True)

    p2 = p()
    p2.space_after = Pt(4)
    add_run(p2, "Compl.　", Pt(12))
    add_run(p2, "26.3%", Pt(24), bold=True)

    p3 = p()
    p3.space_after = Pt(12)
    p3.line_spacing = 1.15
    add_run(p3, "复核第二趟：44.7% / 25.9%（Acc 噪声约 ±3.6pt）", Pt(12))

    p4 = p()
    p4.space_after = Pt(6)
    add_run(p4, "② 能跑到 day 8 的结果", Pt(15), bold=True)

    p5 = p()
    p5.space_after = Pt(2)
    add_run(p5, "干净窗 day01–07　　", Pt(12))
    add_run(p5, "+7.1 pt", Pt(24), bold=True)

    for txt, sp in [
        ("Acc 67.1% vs 基线同窗 60.0%（约四分之一题量）", 4),
        ("day08 起开始被超长回复 / timeout 打穿", 2),
        ("配置：Qwen3-4B-Thinking-2507　｜　30 天 / 346 题　｜　按天会话", 8),
    ]:
        px = p()
        px.space_before = Pt(sp)
        px.space_after = Pt(2)
        px.line_spacing = 1.15
        add_run(px, txt, Pt(11.5))

    right = slide.shapes.add_textbox(Inches(7.00), Inches(1.30), Inches(5.80), Inches(5.80))
    set_rich_text(right.text_frame, [
        (0, "**③ 后续会崩溃**"),
        (1, "day08–21 Acc **17.1%**，基线同窗 48.1%（−31pt）"),
        (1, "复读率：step≤9 为 **0%** → step10 起病 **17%** → step15 **85%**"),
        (1, "崩坏内容是字面重复 thinking（一句话复读两三千次），不是更细的推理"),
        (0, "**④ 还没完成**"),
        (1, "长度爆炸是训练造成的——基线同期没有大面积 timeout"),
        (1, "尚未稳定训满 30 天，不能把干净窗数字当成完整结论"),
    ], Pt(15), Pt(13), line_spacing=1.14, space_after=Pt(7))
    disable_proofing(slide)


def main():
    prs = Presentation(str(PPT))
    n = len(prs.slides)
    assert n == 5, f"expected restored 5-page original, got {n}"

    before = snapshot_slides(prs, 3)
    render_ours_table()
    assert PAPER_TBL.exists(), PAPER_TBL

    update_slide4(prs.slides[3])

    layout = prs.slides[2].slide_layout
    new_slide = prs.slides.add_slide(layout)
    fill_slide5(new_slide, donor=prs.slides[2])
    move_slide(prs, old_index=5, new_index=4)

    after = snapshot_slides(prs, 3)
    if before != after:
        raise RuntimeError("前三页 XML 变了，已中止保存")

    last_text = []
    last = prs.slides[-1]
    for sh in last.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            for s in sh.shapes:
                if s.has_text_frame:
                    last_text.append(s.text_frame.text)
        elif sh.has_text_frame:
            last_text.append(sh.text_frame.text)
    if "THANKS" not in " ".join(last_text):
        raise RuntimeError("最后一页不是 THANKS")

    tmp = HERE / "_share_rebuild_tmp.pptx"
    prs.save(str(tmp))
    try:
        import shutil
        shutil.copy2(tmp, PPT)
        tmp.unlink(missing_ok=True)
        print("saved", PPT.name, "slides", len(prs.slides))
    except PermissionError:
        print("PPT locked, left temp at", tmp)
        print("saved TEMP", tmp.name, "slides", len(prs.slides))
    for i, s in enumerate(prs.slides, 1):
        t = s.shapes.title.text_frame.text if s.shapes.title else "(no title)"
        print(f"  {i}. {t}")


if __name__ == "__main__":
    main()
