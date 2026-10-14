# -*- coding: utf-8 -*-
"""In-place patch slide 4 K=0 numbers. Do not touch slides 1-3 or THANKS."""
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Emu

import rebuild_share_v2 as m

HERE = Path(r"D:\MAO\Claude\openclaw-rl\report")
PPT = HERE / "Openclaw-RL复现及迁移内容分享.pptx"
BAK = HERE / "Openclaw-RL复现及迁移内容分享_backup_更新实验页前_20260917_1001.pptx"


def snapshot(prs, idxs):
    return [etree.tostring(prs.slides[i]._element) for i in idxs]


def main():
    m.render_ours_table()
    prs = Presentation(str(PPT))
    assert len(prs.slides) == 6
    before13 = snapshot(prs, [0, 1, 2])
    before_thanks = snapshot(prs, [5])[0]

    slide = prs.slides[3]
    replaced = 0
    for shape in list(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and 2.7 < shape.top.inches < 3.3:
            left, top, width = shape.left, shape.top, shape.width
            shape._element.getparent().remove(shape._element)
            slide.shapes.add_picture(str(m.OURS_TBL), left, top, width=width)
            replaced += 1
    assert replaced == 1, replaced

    repls = {
        "+2.4pt": "+2.9pt",
        "+0.5pt": "+1.8pt",
        "相对同路径 K=0：": "相对同路径 K=0（09-05）：",
        "K=6 相对同路径 K=0 有提升": "K=0 为 09-05 环境修复后定版；K=6 相对其有提升",
    }
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                t = run.text
                nt = t
                for a, b in repls.items():
                    nt = nt.replace(a, b)
                if nt != t:
                    run.text = nt
                    print("patched run:", repr(nt[:80]))

    m.disable_proofing(slide)
    after13 = snapshot(prs, [0, 1, 2])
    after_thanks = snapshot(prs, [5])[0]
    if before13 != after13:
        raise RuntimeError("前三页 XML 变了")
    if before_thanks != after_thanks:
        raise RuntimeError("THANKS XML 变了")

    tmp = HERE / "_share_k0_tmp.pptx"
    prs.save(str(tmp))
    import shutil
    try:
        shutil.copy2(tmp, PPT)
        tmp.unlink(missing_ok=True)
        print("saved", PPT.name)
    except PermissionError:
        print("locked, left", tmp)


if __name__ == "__main__":
    main()
