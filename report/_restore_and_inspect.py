# -*- coding: utf-8 -*-
import shutil
from pathlib import Path
from pptx import Presentation

HERE = Path(r"D:\MAO\Claude\openclaw-rl\report")
PPT = HERE / "Openclaw-RL复现及迁移内容分享.pptx"
ORIG = HERE / "Openclaw-RL复现及迁移内容分享_backup_更新实验页前_20260917_1001.pptx"
WRONG = HERE / "Openclaw-RL复现及迁移内容分享_backup_错误口径6页_20260917.pptx"

assert ORIG.exists(), ORIG
# 先备份当前错误口径
shutil.copy2(PPT, WRONG)
print("saved wrong-version backup", WRONG.name, WRONG.stat().st_size)

# 退回原版
shutil.copy2(ORIG, PPT)
print("restored original over PPT")

prs = Presentation(str(PPT))
print("SLIDE COUNT", len(prs.slides))
for i, s in enumerate(prs.slides, 1):
    t = s.shapes.title.text_frame.text.strip() if s.shapes.title else "(no title)"
    print(f"  {i}. {t}")
    if i == 4:
        for sh in s.shapes:
            if sh.has_text_frame:
                txt = " | ".join(p.text for p in sh.text_frame.paragraphs if p.text.strip())
                if txt:
                    print("   ", txt[:120])
