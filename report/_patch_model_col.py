# -*- coding: utf-8 -*-
"""Replace slide-4 per-question table PNG only. Do not touch other slides."""
from __future__ import annotations

import shutil
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import rebuild_share_v2 as m

HERE = m.HERE
PPT = m.PPT
BAK = HERE / "Openclaw-RL复现及迁移内容分享_backup_更新实验页前_20260917_1001.pptx"


def snapshot(prs, idxs):
    return [etree.tostring(prs.slides[i]._element) for i in idxs]


def main():
    m.render_ours_table()
    prs = Presentation(str(PPT))
    assert len(prs.slides) == 6
    before = snapshot(prs, [0, 1, 2, 4, 5])

    slide = prs.slides[3]
    replaced = 0
    for shape in list(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and 2.7 < shape.top.inches < 3.3:
            left, top, width = shape.left, shape.top, shape.width
            shape._element.getparent().remove(shape._element)
            slide.shapes.add_picture(str(m.OURS_TBL), left, top, width=width)
            replaced += 1
    assert replaced == 1, replaced
    m.disable_proofing(slide)

    after = snapshot(prs, [0, 1, 2, 4, 5])
    if before != after:
        raise RuntimeError("非第 4 页 XML 变了")

    bak = Presentation(str(BAK))
    for i in range(3):
        if etree.tostring(prs.slides[i]._element) != etree.tostring(bak.slides[i]._element):
            raise RuntimeError(f"第 {i+1} 页相对原版备份变了")

    tmp = HERE / "_share_model_tmp.pptx"
    prs.save(str(tmp))
    try:
        shutil.copy2(tmp, PPT)
        tmp.unlink(missing_ok=True)
        print("saved", PPT.name)
    except PermissionError:
        alt = HERE / "_share_model_ready.pptx"
        shutil.copy2(tmp, alt)
        print("locked, wrote", alt.name)


if __name__ == "__main__":
    main()
