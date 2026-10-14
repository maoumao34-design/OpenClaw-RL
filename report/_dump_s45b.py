# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pathlib import Path

ppt = Path(r"D:\MAO\Claude\openclaw-rl\report\Openclaw-RL复现及迁移内容分享.pptx")
out = Path(r"D:\MAO\Claude\openclaw-rl\report\_dump_s45.txt")
prs = Presentation(str(ppt))
lines = []
for i in (3, 4):
    s = prs.slides[i]
    lines.append(f"===== SLIDE {i+1} =====")
    for sh in s.shapes:
        t = ""
        if sh.has_text_frame:
            t = sh.text_frame.text.replace("\n", " | ")
        lines.append(
            f"  [{sh.name}] L={sh.left.inches:.2f} T={sh.top.inches:.2f} "
            f"W={sh.width.inches:.2f} H={sh.height.inches:.2f} {t}"
        )
out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out)
