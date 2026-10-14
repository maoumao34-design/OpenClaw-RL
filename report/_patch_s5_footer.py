# -*- coding: utf-8 -*-
"""Add full repeated-sentence footer on slide 5. Do not touch table pictures."""
from __future__ import annotations

import shutil
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

import rebuild_share_v2 as m

HERE = Path(r"D:\MAO\Claude\openclaw-rl\report")
PPT = HERE / "Openclaw-RL复现及迁移内容分享.pptx"
QUOTE = (
    "The user's message is a user instruction, "
    "so the assistant's response is a tool call to the write tool."
)


def snapshot(prs, idxs):
    return [etree.tostring(prs.slides[i]._element) for i in idxs]


def fill_conclusion(body):
    body.left, body.top = Inches(6.60), Inches(4.62)
    body.width, body.height = Inches(6.13), Inches(2.20)
    tf = body.text_frame
    tf.clear()
    tf.word_wrap = True

    def conc_p(first=False):
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        para.level = 0
        m.set_no_bullet(para)
        para.line_spacing = 1.12
        para.space_before = Pt(0)
        return para

    p0 = conc_p(True)
    p0.space_after = Pt(6)
    m.add_run(p0, "结论", Pt(18), bold=True)

    p1 = conc_p()
    p1.space_after = Pt(5)
    m.add_run(p1, "① 前八天 Acc 升、Compl 持平；day09–20 相对基线 Acc/Compl 都掉（+5.5pt → −33pt）。", Pt(14))

    p2 = conc_p()
    p2.space_after = Pt(5)
    m.add_run(p2, "② 后半崩，是因为 day09 起大部分题目 timeout。", Pt(14))

    p3 = conc_p()
    p3.space_after = Pt(0)
    m.add_run(p3, "③ timeout 基本是重复 thinking 把 context 顶满。day09 一次 MAIN turn：同一句 ×1388，thinking 15.9 万字，finish_reason=length。", Pt(14))


def add_footer(slide):
    # 对齐四行对照表左缘 L=0.53，拉满到与逐日表同宽
    box = slide.shapes.add_textbox(Inches(0.53), Inches(6.98), Inches(12.20), Inches(0.40))
    tf = box.text_frame
    tf.word_wrap = False
    tf.clear()
    p = tf.paragraphs[0]
    m.set_no_bullet(p)
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    m.add_run(p, QUOTE + "   ", Pt(13))
    m.add_run(p, "×1388", Pt(15), bold=True, color=m.HILITE)
    return box


def main():
    prs = Presentation(str(PPT))
    assert len(prs.slides) == 6
    before = snapshot(prs, [0, 1, 2, 3, 5])
    s5 = prs.slides[4]
    pics = [(sh.name, int(sh.left), int(sh.top), int(sh.width), int(sh.height))
            for sh in s5.shapes if sh.name.startswith("Picture")]

    # 去掉旧底栏（若重跑）
    for sh in list(s5.shapes):
        if sh.has_text_frame and not sh.is_placeholder:
            t = sh.text_frame.text
            if "user's message is a user instruction" in t or "×1388" in t and "timeout" not in t:
                sh._element.getparent().remove(sh._element)

    body = next(sh for sh in s5.placeholders if sh.placeholder_format.idx == 1)
    fill_conclusion(body)
    footer = add_footer(s5)
    m.disable_proofing(s5)

    after = snapshot(prs, [0, 1, 2, 3, 5])
    if before != after:
        raise RuntimeError("非第 5 页 XML 变了")
    pics2 = [(sh.name, int(sh.left), int(sh.top), int(sh.width), int(sh.height))
             for sh in s5.shapes if sh.name.startswith("Picture")]
    if pics != pics2:
        raise RuntimeError("图片被改动")

    tmp = HERE / "_share_s5_tmp.pptx"
    prs.save(str(tmp))
    try:
        shutil.copy2(tmp, PPT)
        tmp.unlink(missing_ok=True)
        print("saved", PPT.name)
    except PermissionError:
        alt = HERE / "_share_s5_ready.pptx"
        shutil.copy2(tmp, alt)
        print("locked, wrote", alt.name)
        return

    prs = Presentation(str(PPT))
    s5 = prs.slides[4]
    print("===== SLIDE 5 =====")
    pair = next(sh for sh in s5.shapes if sh.name.startswith("Picture") and abs(sh.width.inches - 5.90) < 0.05)
    for sh in s5.shapes:
        t = sh.text_frame.text.replace("\n", " | ") if sh.has_text_frame else ""
        print(f"{sh.name:28} L={sh.left.inches:.3f} T={sh.top.inches:.3f} W={sh.width.inches:.3f} H={sh.height.inches:.3f}")
        if t.strip():
            print(" ", t)
    print("pair L", round(pair.left.inches, 3))
    print("footer L", round(footer.left.inches, 3) if footer else None)


if __name__ == "__main__":
    main()
