# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pathlib import Path

ppt = Path(r"D:\MAO\Claude\openclaw-rl\report\Openclaw-RL复现及迁移内容分享.pptx")
out = Path(r"D:\MAO\Claude\openclaw-rl\report\_dump_s4.txt")
prs = Presentation(str(ppt))
s = prs.slides[3]
lines = [f"slide4 shapes={len(s.shapes)}"]
for sh in s.shapes:
    t = ""
    if sh.has_text_frame:
        t = sh.text_frame.text.replace("\n", " | ")
    extra = ""
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        extra = f" pic"
    lines.append(
        f"  {sh.name:28} L={sh.left.inches:.2f} T={sh.top.inches:.2f} "
        f"W={sh.width.inches:.2f} H={sh.height.inches:.2f}{extra} {t[:180]}"
    )
out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out)
