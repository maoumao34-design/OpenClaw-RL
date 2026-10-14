# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pathlib import Path

ppt = Path(r"D:\MAO\Claude\openclaw-rl\report\Openclaw-RL复现及迁移内容分享.pptx")
out = Path(r"D:\MAO\Claude\openclaw-rl\report\_dump_k_text.txt")
prs = Presentation(str(ppt))
lines = [f"slides {len(prs.slides)}"]


def walk(shapes, prefix=""):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            walk(sh.shapes, prefix + "  ")
            continue
        bits = []
        if sh.has_text_frame:
            bits.append(sh.text_frame.text)
        if sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    bits.append(cell.text_frame.text)
        t = "\n".join(bits).strip()
        if t:
            one = t.replace("\n", " | ")
            lines.append(f"{prefix}[{sh.name}] {one}")


for i, s in enumerate(prs.slides, 1):
    lines.append(f"===== SLIDE {i} =====")
    walk(s.shapes)

out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out)
