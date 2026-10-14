# -*- coding: utf-8 -*-
"""更新分享 PPT 第 4 页及之后（前三页一字不改）。

依据 work_log / training_config / metaclaw_migration_plan 最新实验：
- 17.8% → 37.3% 已作废，不再出现
- 定版基线：按天 K=0 Acc 41.1% / Compl 26.3%
- 复核：44.7% / 25.9%，Acc 噪声约 ±3.6pt
- 最新训练 20260914_181842：day01–07 +7.1pt，后期被长度爆炸打穿

用法：python update_share_ppt.py
"""
from __future__ import annotations

import copy
import re
import shutil
from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(r"D:\MAO\Claude\openclaw-rl\report")
PPT = HERE / "Openclaw-RL复现及迁移内容分享.pptx"
BASELINE_IMG = HERE / "metaclaw_baseline_card.png"
DEGEN_IMG = HERE / "metaclaw_degen_bar.png"

EA_FONT = "微软雅黑"
LATIN_FONT = "Segoe UI"
F = "C:/Windows/Fonts/msyh.ttc"
FB = "C:/Windows/Fonts/msyhbd.ttc"

INK = (34, 38, 46)
MUTED = (110, 118, 130)
RED = (176, 42, 42)
RULE = (210, 214, 220)
BAND = (255, 244, 214)
BG = (255, 255, 255)


def style_run(run, size, bold=None, color=None):
    font = run.font
    font.size = size
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = RGBColor(*color)
    font.name = LATIN_FONT
    rPr = font._rPr
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        latin = rPr.find(qn("a:latin"))
        latin.addnext(ea) if latin is not None else rPr.append(ea)
    ea.set("typeface", EA_FONT)
    rPr.set("lang", "zh-CN")
    rPr.set("altLang", "en-US")
    rPr.set("noProof", "1")


def add_run(para, text, size, bold=None, color=None):
    run = para.add_run()
    run.text = text
    style_run(run, size, bold=bold, color=color)
    return run


def reset_bullet(para):
    pPr = para._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    for attr in ("marL", "indent"):
        pPr.attrib.pop(attr, None)


def set_rich_text(text_frame, items, head_size, body_size,
                  line_spacing=1.03, space_after=Pt(4)):
    text_frame.clear()
    text_frame.word_wrap = True
    for i, (level, raw) in enumerate(items):
        para = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        para.level = level
        reset_bullet(para)
        para.line_spacing = line_spacing
        para.space_before = Pt(0)
        para.space_after = space_after if level else Pt(4)
        size = head_size if level == 0 else body_size
        for seg in re.split(r"(\*\*.+?\*\*)", raw):
            if not seg:
                continue
            run = para.add_run()
            bold = seg.startswith("**") and seg.endswith("**")
            run.text = seg[2:-2] if bold else seg
            style_run(run, size, bold=True if bold else None)


def disable_proofing(slide):
    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes)
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    rPr = run.font._rPr
                    rPr.set("lang", "zh-CN")
                    rPr.set("altLang", "en-US")
                    rPr.set("noProof", "1")
    walk(slide.shapes)


def render_baseline_card():
    """左上对照卡：定版基线数字，风格靠近第 3 页结论块。"""
    W, H = 1140, 500
    img = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(img)
    title = ImageFont.truetype(FB, 36)
    label = ImageFont.truetype(F, 32)
    num = ImageFont.truetype(FB, 72)
    note = ImageFont.truetype(F, 26)

    d.rectangle([0, 0, W - 1, H - 1], outline=RULE, width=2)
    d.text((40, 28), "对照基准  ·  零训练（按天 K=0）", font=title, fill=INK)
    d.line([(40, 84), (W - 40, 84)], fill=RULE, width=2)

    rows = [
        ("Acc.", "41.1%"),
        ("Compl.", "26.3%"),
    ]
    y = 120
    for lab, val in rows:
        d.text((48, y + 18), lab, font=label, fill=MUTED)
        tw = d.textbbox((0, 0), val, font=num)[2]
        d.text((W - 56 - tw, y), val, font=num, fill=INK)
        y += 110
        d.line([(40, y - 12), (W - 40, y - 12)], fill=RULE, width=1)

    d.rectangle([24, y + 8, W - 24, H - 24], fill=BAND)
    d.text((48, y + 28), "同配置复核　　44.7%  /  25.9%", font=note, fill=INK)
    d.text((48, y + 72), "Acc 跨趟差约 ±3.6 pt，低于此与噪声不可区分", font=note, fill=MUTED)
    img.save(BASELINE_IMG)
    print(f"wrote {BASELINE_IMG.name} {img.size}")


def render_degen_bar():
    """第 5 页底栏：复读率 0% → 17% → 85%。"""
    W, H = 2400, 220
    img = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(img)
    title = ImageFont.truetype(FB, 36)
    lab = ImageFont.truetype(F, 28)
    num = ImageFont.truetype(FB, 44)
    d.text((40, 16), "训练如何放大退化（复读率）", font=title, fill=INK)

    stages = [
        (0.08, "step ≤9", "0%", "健康", (22, 130, 86)),
        (0.40, "step 10", "17%", "起病", (196, 120, 28)),
        (0.72, "step 15", "85%", "整趟崩溃", RED),
    ]
    y = 120
    d.line([(120, y + 22), (W - 120, y + 22)], fill=RULE, width=6)
    for x_frac, step, pct, name, color in stages:
        x = int(W * x_frac) + 80
        d.ellipse([x - 14, y + 8, x + 14, y + 36], fill=color)
        tw = d.textbbox((0, 0), pct, font=num)[2]
        d.text((x - tw / 2, y - 52), pct, font=num, fill=color)
        sw = d.textbbox((0, 0), f"{step}  {name}", font=lab)[2]
        d.text((x - sw / 2, y + 44), f"{step}  {name}", font=lab, fill=MUTED)
    img.save(DEGEN_IMG)
    print(f"wrote {DEGEN_IMG.name} {img.size}")


def copy_page_number(src_slide, dst_slide):
    existing = [
        sh for sh in dst_slide.placeholders
        if sh.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER
    ]
    if existing:
        return
    donor = next(
        (sh for sh in src_slide.placeholders
         if sh.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER),
        None,
    )
    if donor is None:
        return
    dst_slide.shapes._spTree.append(copy.deepcopy(donor._element))


def move_slide(prs, old_index, new_index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    el = slides[old_index]
    xml_slides.remove(el)
    xml_slides.insert(new_index, el)


def snapshot_slides(prs, n=3):
    return [etree.tostring(prs.slides[i]._element) for i in range(n)]


def update_slide4(slide):
    """保留右栏迁移图与标题；替换左上作废 Table 1；改左下文字。"""
    title = next(sh for sh in slide.placeholders if sh.placeholder_format.idx == 0)
    assert title.text_frame.paragraphs[0].runs[0].text == "方法泛化方案"

    # 只删左上小图（旧 Table 1），右栏大图 left>6 保留
    for shape in list(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.left.inches < 6:
            shape._element.getparent().remove(shape._element)

    # 左上新对照卡，位置尽量贴合原 Table 1
    img_left, img_top, img_w = 0.85, 1.12, 4.78
    pic = slide.shapes.add_picture(
        str(BASELINE_IMG), Inches(img_left), Inches(img_top), width=Inches(img_w)
    )
    img_h = pic.height.inches

    # 更新说明小字
    caption = None
    for shape in slide.shapes:
        if shape.has_text_frame and not shape.is_placeholder:
            t = "".join(p.text for p in shape.text_frame.paragraphs)
            if "仅迁移 RL" in t or "数量级参照" in t or "Table 1" in t:
                caption = shape
                break
    cap_text = "不含 MetaClaw 技能库；训练效果只跟自己的零训练基线比，不直接对照论文 Table 1"
    cap_left, cap_w = 0.53, 5.74
    if caption is None:
        caption = slide.shapes.add_textbox(
            Inches(cap_left), Inches(img_top + img_h + 0.08),
            Inches(cap_w), Inches(0.42),
        )
    else:
        caption.left = Inches(cap_left)
        caption.top = Inches(img_top + img_h + 0.06)
        caption.width = Inches(cap_w)
        caption.height = Inches(0.42)
    tf = caption.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.line_spacing = 1.08
    add_run(p, cap_text, Pt(10.0))

    body = next(sh for sh in slide.placeholders if sh.placeholder_format.idx == 1)
    text_top = img_top + img_h + 0.56
    body.left, body.width = Inches(0.53), Inches(5.74)
    body.top = Inches(text_top)
    body.height = Inches(max(1.70, 7.32 - text_top))
    set_rich_text(body.text_frame, [
        (0, "**① 任务背景**"),
        (1, "方法已在 Personal Agent 场景复现校准，换到 **MetaClaw-Bench**"
            "（30 天连续任务流、CLI agent）验证方法能否迁移"),
        (0, "**② 当前进展**"),
        (1, "零训练基线已定版；训练前半段相对基线有提升"),
        (1, "后期会被长度爆炸打穿，尚未稳定跑满 30 天"),
    ], Pt(13.5), Pt(11.5), line_spacing=1.03, space_after=Pt(4))
    disable_proofing(slide)


def fill_slide5(slide, donor):
    """新增：迁移实验结果。"""
    copy_page_number(donor, slide)

    title = next(sh for sh in slide.placeholders if sh.placeholder_format.idx == 0)
    title.left, title.top = Inches(0.53), Inches(0.28)
    title.width, title.height = Inches(9.49), Inches(0.90)
    tf = title.text_frame
    tf.clear()
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    add_run(tf.paragraphs[0], "迁移实验结果", Pt(26), bold=True)

    # 正文占位符改成左栏数字块；底部留给退化条
    body = next(sh for sh in slide.placeholders if sh.placeholder_format.idx == 1)
    body.left, body.top = Inches(0.53), Inches(1.22)
    body.width, body.height = Inches(6.20), Inches(4.55)
    tf = body.text_frame
    tf.clear()
    tf.word_wrap = True

    def p(first=False):
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        para.level = 0
        pPr = para._p.get_or_add_pPr()
        pPr.attrib["marL"] = "0"
        pPr.attrib["indent"] = "0"
        if not pPr.findall(qn("a:buNone")):
            pPr.append(pPr.makeelement(qn("a:buNone"), {}))
        return para

    p0 = p(True)
    p0.space_after = Pt(6)
    add_run(p0, "零训练基线（按天 K=0）", Pt(15), bold=True)

    p1 = p()
    p1.space_after = Pt(2)
    add_run(p1, "Acc.　　", Pt(12))
    add_run(p1, "41.1%", Pt(24), bold=True)

    p2 = p()
    p2.space_after = Pt(4)
    add_run(p2, "Compl.　", Pt(12))
    add_run(p2, "26.3%", Pt(24), bold=True)

    p3 = p()
    p3.space_before = Pt(2)
    p3.space_after = Pt(10)
    p3.line_spacing = 1.15
    add_run(p3, "同配置再跑一趟：44.7% / 25.9%（Acc 噪声约 ±3.6pt）", Pt(11))

    p4 = p()
    p4.space_after = Pt(6)
    add_run(p4, "最新训练 · 干净窗（day01–07）", Pt(15), bold=True)

    p5 = p()
    p5.space_after = Pt(2)
    add_run(p5, "相对基线同窗　　", Pt(12))
    add_run(p5, "+7.1 pt", Pt(24), bold=True)

    for txt, sp_before in [
        ("Acc 67.1% vs 基线同窗 60.0%（约四分之一题量，还不是完整结论）", 6),
        ("day08 起被超长回复打穿，全程未跑满 30 天", 2),
        ("配置：Policy = Qwen3-4B-Thinking-2507　｜　30 天 / 346 题　｜　按天会话", 8),
    ]:
        px = p()
        px.space_before = Pt(sp_before)
        px.space_after = Pt(2)
        px.line_spacing = 1.15
        add_run(px, txt, Pt(11))

    # 右栏：已确认 / 卡在哪
    right = slide.shapes.add_textbox(Inches(7.00), Inches(1.22), Inches(5.80), Inches(4.55))
    set_rich_text(right.text_frame, [
        (0, "**① 已确认**"),
        (1, "长度爆炸是训练造成的——基线同期没有大面积 timeout"),
        (1, "训练会放大退化：一次 17% 的复读扰动，被后续更新打成 85% 的整趟崩溃"),
        (0, "**② 卡在哪**"),
        (1, "还没有一趟「稳定训满 30 天、全程显著高于基线」的完整结果"),
        (1, "下一步：拦住退化样本后再训，看曲线是否走平"),
        (0, "**③ 对照口径**"),
        (1, "全套工具 agent + 30 天 / 346 题，跟论文 Table 1 任何一格都不是同一条件，不能直接比"),
    ], Pt(15), Pt(13), line_spacing=1.12, space_after=Pt(7))

    slide.shapes.add_picture(
        str(DEGEN_IMG), Inches(0.53), Inches(5.90), width=Inches(12.27)
    )
    disable_proofing(slide)


def latest_backup():
    cands = sorted(HERE.glob("Openclaw-RL复现及迁移内容分享_backup_更新实验页前_*.pptx"))
    return cands[-1] if cands else None


def main():
    assert PPT.exists(), PPT
    bak = latest_backup()
    if bak is not None:
        shutil.copy2(bak, PPT)
        print(f"restore from {bak.name}")
    else:
        stamp = datetime.now().strftime("%Y%m%d_%H%M")
        bak = HERE / f"Openclaw-RL复现及迁移内容分享_backup_更新实验页前_{stamp}.pptx"
        shutil.copy2(PPT, bak)
        print(f"backup → {bak.name}")

    render_baseline_card()
    render_degen_bar()
    prs = Presentation(str(PPT))
    assert len(prs.slides) == 5, f"unexpected slide count {len(prs.slides)}"

    before = snapshot_slides(prs, 3)
    titles0 = [
        "".join(r.text for r in prs.slides[i].shapes.title.text_frame.paragraphs[0].runs)
        if prs.slides[i].shapes.title else ""
        for i in range(3)
    ]

    update_slide4(prs.slides[3])

    layout = prs.slides[2].slide_layout  # Title and Content，与第 3 页相同
    new_slide = prs.slides.add_slide(layout)
    fill_slide5(new_slide, donor=prs.slides[2])
    move_slide(prs, old_index=5, new_index=4)

    after = snapshot_slides(prs, 3)
    if before != after:
        raise RuntimeError("前三页 XML 发生了变化，已中止保存")
    titles1 = [
        "".join(r.text for r in prs.slides[i].shapes.title.text_frame.paragraphs[0].runs)
        if prs.slides[i].shapes.title else ""
        for i in range(3)
    ]
    if titles0 != titles1:
        raise RuntimeError(f"前三页标题变了：{titles0} → {titles1}")

    # THANKS 必须仍是最后一页
    last = prs.slides[-1]
    last_text = []
    for sh in last.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            for s in sh.shapes:
                if s.has_text_frame:
                    last_text.append(s.text_frame.text)
        elif sh.has_text_frame:
            last_text.append(sh.text_frame.text)
    blob = " ".join(last_text)
    if "THANKS" not in blob:
        raise RuntimeError(f"最后一页不是 THANKS：{blob!r}")

    prs.save(str(PPT))
    print(f"saved {PPT.name}  slides={len(prs.slides)}")
    for i, s in enumerate(prs.slides, 1):
        t = s.shapes.title.text_frame.text if s.shapes.title else "(no title)"
        print(f"  {i}. {t}")
    print(f"backup: {bak}")


if __name__ == "__main__":
    main()
