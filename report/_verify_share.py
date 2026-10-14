# -*- coding: utf-8 -*-
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

HERE = Path(r"D:\MAO\Claude\openclaw-rl\report")
OUT = HERE / "_verify_share.txt"
PPT = HERE / "Openclaw-RL复现及迁移内容分享.pptx"
BAK = HERE / "Openclaw-RL复现及迁移内容分享_backup_更新实验页前_20260917_1001.pptx"


def dump_slide(prs, i, lines):
    s = prs.slides[i]
    title = ""
    try:
        title = s.shapes.title.text_frame.text.strip().splitlines()[0] if s.shapes.title else ""
    except Exception:
        pass
    lines.append(f"-- slide {i+1} {title!r} shapes={len(s.shapes)}")
    for sh in s.shapes:
        t = ""
        if sh.has_text_frame:
            t = sh.text_frame.text.replace("\n", " | ")[:240]
        lines.append(
            f"  {sh.name:28} L={sh.left.inches:.2f} T={sh.top.inches:.2f} "
            f"W={sh.width.inches:.2f} H={sh.height.inches:.2f} {t}"
        )


def main():
    lines = []
    pb = Presentation(str(BAK))
    pc = Presentation(str(PPT))
    lines.append(f"backup slides={len(pb.slides)} current slides={len(pc.slides)}")
    for i in range(3):
        a = etree.tostring(pb.slides[i]._element)
        b = etree.tostring(pc.slides[i]._element)
        lines.append(f"slide {i+1} xml == backup: {a == b}")
    last_b = etree.tostring(pb.slides[-1]._element)
    last_c = etree.tostring(pc.slides[-1]._element)
    lines.append(f"THANKS xml == backup last: {last_b == last_c}")
    lines.append("=== CURRENT 4-6 ===")
    for i in range(3, len(pc.slides)):
        dump_slide(pc, i, lines)
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print("wrote", OUT)


if __name__ == "__main__":
    main()
