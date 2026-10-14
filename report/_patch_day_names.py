# -*- coding: utf-8 -*-
"""Rename experiment K=0/K=6 to day 0 / day 6 on slides 4-5 only."""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

import rebuild_share_v2 as m

HERE = Path(r"D:\MAO\Claude\openclaw-rl\report")
PPT = HERE / "Openclaw-RL复现及迁移内容分享.pptx"
DUMP = HERE / "_dump_k_text.txt"

REPLACEMENTS = [
    ("K=6", "day 6"),
    ("k=6", "day 6"),
    ("K＝6", "day 6"),
    ("K=0", "day 0"),
    ("k=0", "day 0"),
    ("K＝0", "day 0"),
]


def snapshot(prs, idxs):
    return [etree.tostring(prs.slides[i]._element) for i in idxs]


def apply_text(s: str) -> str:
    out = s
    for a, b in REPLACEMENTS:
        out = out.replace(a, b)
    return out


def patch_runs(shape, hits):
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                nt = apply_text(run.text)
                if nt != run.text:
                    hits.append(nt)
                    run.text = nt
    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                patch_runs(cell, hits)


def walk_shapes(shapes, hits):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            walk_shapes(sh.shapes, hits)
        else:
            patch_runs(sh, hits)


def collect_text(slide):
    lines = []

    def walk(shapes, prefix=""):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(sh.shapes, prefix + "  ")
                continue
            bits = []
            if sh.has_text_frame:
                bits.append(sh.text_frame.text)
            if sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        bits.append(cell.text_frame.text)
            t = "\n".join(bits).strip()
            if t:
                lines.append(f"{prefix}[{sh.name}] {t.replace(chr(10), ' | ')}")

    walk(slide.shapes)
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            lines.append(f"[NOTES] {notes.replace(chr(10), ' | ')}")
    return lines


def remaining_k(text: str):
    return re.findall(r"[Kk][=＝]\s*[06]|K=0|K=6", text)


def main():
    m.render_ours_table()
    prs = Presentation(str(PPT))
    assert len(prs.slides) == 6
    before13 = snapshot(prs, [0, 1, 2])
    before_thanks = snapshot(prs, [5])[0]

    hits = []
    for idx in (3, 4):
        slide = prs.slides[idx]
        walk_shapes(slide.shapes, hits)
        if slide.has_notes_slide:
            tf = slide.notes_slide.notes_text_frame
            for para in tf.paragraphs:
                for run in para.runs:
                    nt = apply_text(run.text)
                    if nt != run.text:
                        hits.append(nt)
                        run.text = nt

    slide4 = prs.slides[3]
    replaced = 0
    for shape in list(slide4.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and 2.7 < shape.top.inches < 3.3:
            left, top, width = shape.left, shape.top, shape.width
            shape._element.getparent().remove(shape._element)
            slide4.shapes.add_picture(str(m.OURS_TBL), left, top, width=width)
            replaced += 1
    assert replaced == 1, replaced
    m.disable_proofing(slide4)
    m.disable_proofing(prs.slides[4])

    after13 = snapshot(prs, [0, 1, 2])
    after_thanks = snapshot(prs, [5])[0]
    if before13 != after13:
        raise RuntimeError("前三页 XML 变了")
    if before_thanks != after_thanks:
        raise RuntimeError("THANKS XML 变了")

    tmp = HERE / "_share_day_tmp.pptx"
    prs.save(str(tmp))
    try:
        shutil.copy2(tmp, PPT)
        tmp.unlink(missing_ok=True)
    except PermissionError:
        print("locked, left", tmp)
        return

    prs = Presentation(str(PPT))
    dump = [f"slides {len(prs.slides)}"]
    leftover = []
    for i in (3, 4):
        dump.append(f"===== SLIDE {i+1} =====")
        texts = collect_text(prs.slides[i])
        dump.extend(texts)
        blob = "\n".join(texts)
        leftover.extend([(i + 1, m_) for m_ in remaining_k(blob)])
        xml = etree.tostring(prs.slides[i]._element, encoding="unicode")
        leftover.extend([(i + 1, m_) for m_ in remaining_k(xml)])

    DUMP.write_text("\n".join(dump), encoding="utf-8")
    print("HITS")
    for h in hits:
        print(" ", h)
    print("LEFTOVER", leftover)
    print("DUMP\n" + "\n".join(dump))


if __name__ == "__main__":
    main()
